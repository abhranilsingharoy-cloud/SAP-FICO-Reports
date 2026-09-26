from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Calibri'

def add_paragraph(doc, text, bold_words=None):
    p = doc.add_paragraph()
    p.style.font.name = 'Calibri'
    p.style.font.size = Pt(11)
    
    if bold_words:
        parts = text.split(bold_words)
        if len(parts) > 1:
            p.add_run(parts[0])
            p.add_run(bold_words).bold = True
            p.add_run(parts[1])
        else:
            p.add_run(text)
    else:
        p.add_run(text)
    return p

doc = Document()

# Title
title = doc.add_heading('SAP FICO Reports: Comprehensive Overview, Use Cases, and Benefits', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# Introduction
add_heading(doc, '1. Introduction to SAP FICO', 1)
doc.add_paragraph(
    "SAP FICO (Financial Accounting and Controlling) is the core module of SAP ERP that helps organizations manage their financial data. "
    "The FI (Financial Accounting) component is designed for external reporting, such as generating balance sheets and profit and loss (P&L) statements "
    "for stakeholders, investors, and tax authorities. The CO (Controlling) component is tailored for internal reporting, aiding management in decision-making, "
    "cost control, and profitability analysis. Together, they provide a robust reporting framework that ensures financial transparency, compliance, and strategic alignment."
)

# FI Reports
add_heading(doc, '2. Key Financial Accounting (FI) Reports', 1)

add_heading(doc, '2.1 Balance Sheet and Profit & Loss (P&L) Statement', 2)
doc.add_paragraph(
    "Overview: These are the primary financial statements required for statutory reporting. The Balance Sheet shows the company's assets, liabilities, and equity at a specific point in time, while the P&L statement details income and expenses over a period."
)
doc.add_paragraph("Use Case: Used at the end of a fiscal period (month, quarter, or year) to report financial performance to external stakeholders and regulatory bodies.")
doc.add_paragraph("Benefits: Ensures compliance with accounting standards (e.g., GAAP, IFRS), provides a snapshot of financial health, and aids investors in assessing the company's profitability.")

add_heading(doc, '2.2 General Ledger (GL) Account Balances', 2)
doc.add_paragraph(
    "Overview: This report provides a detailed view of the balances for all general ledger accounts, including opening balances, debits, credits, and closing balances for a specific period."
)
doc.add_paragraph("Use Case: Used by accountants during the month-end closing process to verify that all journal entries have been posted correctly and to prepare the trial balance.")
doc.add_paragraph("Benefits: Facilitates accurate financial reconciliation, helps identify posting errors quickly, and serves as the foundation for the balance sheet and P&L statement.")

add_heading(doc, '2.3 Accounts Receivable (AR) Aging Report', 2)
doc.add_paragraph(
    "Overview: A detailed report categorizing open customer invoices based on the number of days they are past due (e.g., 0-30 days, 31-60 days, 61-90 days, 90+ days)."
)
doc.add_paragraph("Use Case: Used by the credit and collections team to identify overdue accounts, prioritize collection efforts, and assess credit risk.")
doc.add_paragraph("Benefits: Improves cash flow by accelerating collections, reduces bad debt, and provides visibility into customer payment behavior.")

add_heading(doc, '2.4 Accounts Payable (AP) Aging Report', 2)
doc.add_paragraph(
    "Overview: Similar to the AR Aging report, but focuses on unpaid vendor invoices, categorizing them by how long they have been outstanding."
)
doc.add_paragraph("Use Case: Used by the AP department to manage outgoing cash flows, schedule payments, and ensure vendors are paid on time to avoid late fees or leverage early payment discounts.")
doc.add_paragraph("Benefits: Optimizes working capital, maintains good vendor relationships, and ensures accurate cash flow forecasting.")


# CO Reports
add_heading(doc, '3. Key Controlling (CO) Reports', 1)

add_heading(doc, '3.1 Cost Center Actual/Plan/Variance Report', 2)
doc.add_paragraph(
    "Overview: Compares the actual costs incurred by a specific department (cost center) against the planned (budgeted) costs, highlighting any variances."
)
doc.add_paragraph("Use Case: Used by department managers and controllers on a monthly basis to monitor departmental spending and investigate significant overspending or underspending.")
doc.add_paragraph("Benefits: Enables strict cost control, promotes accountability among department heads, and facilitates accurate budgeting for future periods.")

add_heading(doc, '3.2 Profit Center Accounting Report', 2)
doc.add_paragraph(
    "Overview: Evaluates the profitability of individual, independent areas within an organization (profit centers), such as specific product lines, regions, or divisions."
)
doc.add_paragraph("Use Case: Used by senior management to determine which business segments are most profitable and where strategic investments should be directed.")
doc.add_paragraph("Benefits: Supports decentralized decision-making, helps identify underperforming segments, and aligns operational activities with corporate financial goals.")

add_heading(doc, '3.3 Profitability Analysis (CO-PA) Report', 2)
doc.add_paragraph(
    "Overview: Analyzes profitability from a market segment perspective, such as profitability by customer, product, region, or distribution channel. It can be based on costing (Costing-based CO-PA) or general ledger accounts (Account-based CO-PA)."
)
doc.add_paragraph("Use Case: Used by sales and marketing executives to analyze the margins of specific products or customer groups and to determine pricing strategies.")
doc.add_paragraph("Benefits: Provides deep, multidimensional insights into what drives profit, allowing the business to optimize its product mix, adjust pricing, and target the most lucrative market segments.")


# Benefits
add_heading(doc, '4. General Benefits of SAP FICO Reporting', 1)
doc.add_paragraph(
    "1. Real-Time Data Access: Because SAP ERP is highly integrated, FICO reports draw data in real-time from other modules like Sales and Distribution (SD) and Materials Management (MM). This eliminates data lag and ensures reports reflect the current state of the business."
)
doc.add_paragraph(
    "2. High Data Accuracy and Integrity: Automated postings and built-in validations reduce the risk of human error, ensuring that financial reports are reliable and audit-ready."
)
doc.add_paragraph(
    "3. Regulatory Compliance: Standard SAP FI reports are pre-configured to meet country-specific legal and tax requirements, making statutory reporting much simpler for multinational corporations."
)
doc.add_paragraph(
    "4. Enhanced Decision-Making: The detailed, granular insights provided by CO reports empower management to make informed, data-driven decisions regarding cost reduction and revenue growth."
)

# Conclusion
add_heading(doc, '5. Conclusion', 1)
doc.add_paragraph(
    "SAP FICO reporting is an indispensable tool for modern enterprises. While FI reports ensure that the organization remains compliant and transparent to the outside world, CO reports act as a compass for internal management, guiding the company toward greater efficiency and profitability. Mastering these reports allows organizations to maintain a tight grip on their financial health and quickly adapt to changing market dynamics."
)

doc.save('SAP_FICO_Reports.docx')
print("Document generated successfully.")
